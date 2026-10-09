#!/usr/bin/env python3
"""Generic master workbook -> shared ALL-country reference CSV exporter
(maintainer/instructor tool).

Produces the tracked, student-safe data/reference/pollitik_reference.csv:
every country's existing Master-sheet rows, in the same 8 canonical
columns as a per-country countries/<slug>/data/<slug>_reference_snapshot.csv
(see python/build_country_reference_snapshot.py, whose CANON_COLUMNS and
header-matching normalization this script reuses directly rather than
reimplementing). Unlike that script, this one does not filter by country -
it is the single shared source a country workspace can fall back to when
the production master workbook itself isn't present in a checkout.

It never writes to the production master workbook (read-only open, never
a save/write call against it anywhere in this file), never infers a
leader, office, party, question wording, or series alias, and never
exports a column beyond the 8 canonical ones.

Usage:
    python python/build_reference_csv.py --workbook data/master/pollitik_master.xlsx
"""

import argparse
import csv
import json
import os
import sys

import pollitik_common as pc
from build_country_reference_snapshot import (
    CANON_COLUMNS,
    SnapshotBuildError,
    _find_canonical_columns,
    sha256_of_file,
)


def build(workbook_path, output_path=None, sheet_name="Master", overwrite=False):
    output_path = output_path or pc.SHARED_REFERENCE_CSV_PATH

    if not os.path.isfile(workbook_path):
        raise SnapshotBuildError("Master workbook not found at '{}'.".format(workbook_path))

    if os.path.exists(output_path) and not overwrite:
        raise SnapshotBuildError(
            "'{}' already exists - refusing to overwrite it. Pass "
            "--overwrite only if you deliberately intend to replace the "
            "shared reference CSV.".format(output_path)
        )

    workbook_sha256 = sha256_of_file(workbook_path)

    import openpyxl

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise SnapshotBuildError(
                "Sheet '{}' not found in '{}'. Available sheets: {}".format(
                    sheet_name, workbook_path, wb.sheetnames
                )
            )
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            raise SnapshotBuildError("Sheet '{}' has no header row.".format(sheet_name))

        col = _find_canonical_columns(header)

        rows = []
        country_row_counts = {}
        for row in rows_iter:
            country_idx = col["Country"]
            raw_country = row[country_idx] if country_idx < len(row) else None
            if raw_country is None or str(raw_country).strip() == "":
                continue

            out_row = {}
            for c in CANON_COLUMNS:
                idx = col[c]
                v = row[idx] if idx < len(row) else None
                out_row[c] = "" if v is None else str(v)
            rows.append(out_row)

            country_key = pc.normalize_identifier(raw_country)
            country_row_counts[country_key] = country_row_counts.get(country_key, 0) + 1
    finally:
        wb.close()

    if not rows:
        raise SnapshotBuildError(
            "Zero rows with a non-empty Country found in '{}' sheet '{}'. "
            "Refusing to write an empty shared reference CSV - nothing was "
            "written.".format(workbook_path, sheet_name)
        )

    out_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(out_dir, exist_ok=True)
    tmp_path = output_path + ".tmp"
    with open(tmp_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CANON_COLUMNS)
        writer.writeheader()
        for r in rows:
            writer.writerow(r)
    os.replace(tmp_path, output_path)

    return {
        "row_count": len(rows),
        "country_count": len(country_row_counts),
        "country_row_counts": country_row_counts,
        "output_path": output_path,
        "output_size_bytes": os.path.getsize(output_path),
        "workbook_path": workbook_path,
        "workbook_sha256": workbook_sha256,
        "sheet": sheet_name,
        "overwrite_used": bool(overwrite),
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", default=pc.DEFAULT_MASTER_WORKBOOK)
    parser.add_argument("--sheet", default="Master")
    parser.add_argument("--output", default=None, help="Override the tracked shared-CSV output path.")
    parser.add_argument("--overwrite", action="store_true", help="Deliberately overwrite an existing tracked shared CSV.")
    args = parser.parse_args(argv)

    try:
        result = build(args.workbook, output_path=args.output, sheet_name=args.sheet, overwrite=args.overwrite)
    except SnapshotBuildError as exc:
        print(json.dumps({"built": False, "error": str(exc)}, indent=2, ensure_ascii=True))
        return 1

    # ensure_ascii=True: country_row_counts keys can legitimately contain
    # invisible Unicode (NBSP/ZWSP/etc, see build_country_reference_snapshot.py)
    # - printing them literally can crash a cp1252 Windows console.
    print(json.dumps({"built": True, **result}, indent=2, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    sys.exit(main())
