#!/usr/bin/env python3
"""Generic EAD workbook -> country EAD Series/Question-Wording snapshot
exporter (maintainer/instructor tool, student-facing default source).

Produces countries/<slug>/data/<slug>_ead_series_question_wording.csv:
one country's rows from "EAD Series and Question Wording.xlsx", in 6
canonical columns (Country, Series, Description, Question Type, Question
Wording, Method). Mirrors python/build_country_reference_snapshot.py's
shape and precedence exactly:

  1. Prefer reading directly from the real EAD workbook when it's
     present (--ead-workbook, default data/master/EAD Series and
     Question Wording.xlsx).
  2. Otherwise fall back to the shared, tracked ALL-country CSV
     (--shared-ead-csv, default data/reference/ead_series_question_wording.csv,
     built by python/build_ead_reference_csv.py) - this is what lets a
     student workspace with no local Excel workbook still build a real
     per-country EAD snapshot.
  3. Never invents a value, never writes to either source (both opened
     read-only), never merges an alias/formatting-variant Series - that
     judgment stays with report_country_ead_coverage.py / a human.

Country matching uses the shared, deterministic
pollitik_common.normalize_identifier() helper - the same one that
correctly finds the EAD workbook's whitespace-variant Country cells
(e.g. 'Australia ') instead of silently reporting zero rows. The verbose
real "Method (face to face, telephone, mixed, etc.)" header is matched
by prefix ("Method"), the same rule report_country_ead_coverage.py
already uses, then canonicalized to the short name "Method" in the
output - never a reimplemented/ad hoc comparison.

The --country input is first passed through
pollitik_common.resolve_country_input() - the same controlled alias
table build_country_reference_snapshot.py uses (never fuzzy matching),
so e.g. "France President" resolves to the Master value "France_Pres"
here too. A bare "France" is deliberately never auto-resolved (see
AmbiguousCountryError).

Usage:
    python python/build_country_ead_snapshot.py --country "New Zealand"
"""

import argparse
import csv
import json
import os
import sys

import pollitik_common as pc
from build_country_reference_snapshot import sha256_of_file

CANON_COLUMNS = ["Country", "Series", "Description", "Question Type", "Question Wording", "Method"]
_EXACT_COLUMNS = ["Country", "Series", "Description", "Question Type", "Question Wording"]


class EadSnapshotBuildError(Exception):
    """A clearly-reportable data/setup problem. Always caught in main()
    and printed as a single clean JSON error, never an unhandled
    traceback, and never a guessed/invented value."""


def default_output_path(project_dir, slug):
    return os.path.join(project_dir, "countries", slug, "data", "{}_ead_series_question_wording.csv".format(slug))


def _find_canonical_ead_columns(header, source_desc):
    """Maps each of the 6 canonical column names to its index in
    `header`, matching by whitespace-STRIPPED (not casefolded) exact
    text for the first 5, and by a "Method" PREFIX match for the 6th
    (the real header is the long "Method (face to face, telephone,
    mixed, etc.)" string - the same startswith() rule
    report_country_ead_coverage.py already uses). Refuses loudly on any
    missing column rather than guessing at, renaming, or silently
    matching the wrong one."""
    stripped_index = {}
    for i, h in enumerate(header):
        if h is None:
            continue
        key = str(h).strip()
        if key and key not in stripped_index:
            stripped_index[key] = i

    missing = [c for c in _EXACT_COLUMNS if c not in stripped_index]

    method_idx = None
    for key, i in stripped_index.items():
        if key.startswith("Method"):
            method_idx = i
            break
    if method_idx is None:
        missing.append("Method*")

    if missing:
        raise EadSnapshotBuildError(
            "'{}' is missing required column(s): {}. Headers found: {}".format(
                source_desc, ", ".join(missing), [h for h in header if h not in (None, "")]
            )
        )

    col = {c: stripped_index[c] for c in _EXACT_COLUMNS}
    col["Method"] = method_idx
    return col


def _iter_canonical_rows_from_workbook(workbook_path, sheet_name):
    import openpyxl

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise EadSnapshotBuildError(
                "Sheet '{}' not found in '{}'. Available sheets: {}".format(
                    sheet_name, workbook_path, wb.sheetnames
                )
            )
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            raise EadSnapshotBuildError("Sheet '{}' has no header row.".format(sheet_name))

        col = _find_canonical_ead_columns(header, "{} sheet '{}'".format(workbook_path, sheet_name))
        for row in rows_iter:
            out_row = {}
            for c in CANON_COLUMNS:
                idx = col[c]
                v = row[idx] if idx < len(row) else None
                out_row[c] = "" if v is None else str(v)
            yield out_row
    finally:
        wb.close()


def _iter_canonical_rows_from_shared_csv(csv_path):
    with open(csv_path, "r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        header = [h.strip() if h else h for h in (reader.fieldnames or [])]
        missing = [c for c in CANON_COLUMNS if c not in header]
        if missing:
            raise EadSnapshotBuildError(
                "Shared EAD reference CSV '{}' is missing required column(s): "
                "{}. Headers found: {}".format(csv_path, ", ".join(missing), reader.fieldnames)
            )
        for row in reader:
            yield {c: (row.get(c) or "") for c in CANON_COLUMNS}


def build(country, project_dir=None, output_path=None, overwrite=False,
          ead_workbook_path=None, shared_ead_csv_path=None, sheet_name="Sheet1"):
    project_dir = project_dir or pc.PROJECT_DIR

    try:
        resolved_country = pc.resolve_country_input(country)
        slug = pc.canonical_country_slug(country)
    except pc.AmbiguousCountryError as exc:
        raise EadSnapshotBuildError(str(exc))

    country_dir = os.path.join(project_dir, "countries", slug)
    if not os.path.isdir(country_dir):
        raise EadSnapshotBuildError(
            "countries/{}/ does not exist yet - run 'python "
            "python/init_country_workspace.py --country {}' first. This "
            "script never creates a country workspace itself.".format(slug, country)
        )

    output_path = output_path or default_output_path(project_dir, slug)
    if os.path.exists(output_path) and not overwrite:
        raise EadSnapshotBuildError(
            "'{}' already exists - refusing to overwrite it. Pass "
            "--overwrite only if you deliberately intend to replace the "
            "tracked EAD snapshot.".format(output_path)
        )

    ead_workbook_path = ead_workbook_path or pc.EAD_WORDING_REFERENCE_PATH
    shared_ead_csv_path = shared_ead_csv_path or pc.SHARED_EAD_REFERENCE_CSV_PATH

    if os.path.isfile(ead_workbook_path):
        source = "ead_workbook"
        source_path = ead_workbook_path
        source_sha256 = sha256_of_file(ead_workbook_path)
        row_iter = _iter_canonical_rows_from_workbook(ead_workbook_path, sheet_name)
    elif os.path.isfile(shared_ead_csv_path):
        source = "shared_ead_reference_csv"
        source_path = shared_ead_csv_path
        source_sha256 = sha256_of_file(shared_ead_csv_path)
        row_iter = _iter_canonical_rows_from_shared_csv(shared_ead_csv_path)
    else:
        raise EadSnapshotBuildError(
            "Neither the EAD workbook ('{}') nor the shared EAD reference "
            "CSV ('{}') was found. At least one is required to build a "
            "country EAD snapshot - run python/build_ead_reference_csv.py "
            "once (maintainer/instructor step, needs the real workbook) if "
            "only the shared CSV is normally missing.".format(
                ead_workbook_path, shared_ead_csv_path
            )
        )

    target_norm = pc.normalize_identifier(resolved_country)

    matched_rows = []
    raw_variant_counts = {}
    distinct_series = set()

    for out_row in row_iter:
        raw_country = out_row.get("Country")
        if not raw_country:
            continue
        if pc.normalize_identifier(raw_country) != target_norm:
            continue

        raw_variant_counts[raw_country] = raw_variant_counts.get(raw_country, 0) + 1
        matched_rows.append(out_row)

        series_val = out_row.get("Series")
        if series_val:
            distinct_series.add(series_val)

    if not matched_rows:
        raise EadSnapshotBuildError(
            "Zero rows in '{}' ({}) matched Country == '{}' (resolved "
            "from input '{}', after normalization). Refusing to write "
            "an empty snapshot - nothing was written.".format(
                source_path, source, resolved_country, country
            )
        )

    out_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(out_dir, exist_ok=True)
    tmp_path = output_path + ".tmp"
    with open(tmp_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CANON_COLUMNS)
        writer.writeheader()
        for r in matched_rows:
            writer.writerow(r)
    os.replace(tmp_path, output_path)

    return {
        "country": country,
        "resolved_country": resolved_country,
        "country_slug": slug,
        "normalized_lookup_country": target_norm,
        "raw_country_variants_found": raw_variant_counts,
        "normalization_collapsed_multiple_raw_variants": len(raw_variant_counts) > 1,
        "row_count": len(matched_rows),
        "distinct_series_count": len(distinct_series),
        "output_path": output_path,
        "output_size_bytes": os.path.getsize(output_path),
        "sheet": sheet_name,
        "overwrite_used": bool(overwrite),
        "source": source,
        "source_path": source_path,
        "source_sha256": source_sha256,
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--country", required=True, help='e.g. "New Zealand"')
    parser.add_argument("--project-dir", default=None)
    parser.add_argument("--output", default=None, help="Override the tracked EAD snapshot output path.")
    parser.add_argument("--overwrite", action="store_true", help="Deliberately overwrite an existing tracked EAD snapshot.")
    parser.add_argument("--ead-workbook", default=None, help="Override the EAD workbook path (preferred source when present).")
    parser.add_argument(
        "--shared-ead-csv", default=None,
        help="Override the shared all-country EAD CSV path used as a fallback "
             "source when --ead-workbook is not found (default: "
             "data/reference/ead_series_question_wording.csv).",
    )
    parser.add_argument("--sheet", default="Sheet1")
    args = parser.parse_args(argv)

    try:
        result = build(
            args.country, project_dir=args.project_dir, output_path=args.output,
            overwrite=args.overwrite, ead_workbook_path=args.ead_workbook,
            shared_ead_csv_path=args.shared_ead_csv, sheet_name=args.sheet,
        )
    except EadSnapshotBuildError as exc:
        print(json.dumps({"built": False, "error": str(exc)}, indent=2, ensure_ascii=True))
        return 1

    # ensure_ascii=True (matches build_country_reference_snapshot.py):
    # raw_country_variants_found can legitimately contain invisible
    # Unicode - printing it literally can crash a cp1252 Windows console.
    print(json.dumps({"built": True, **result}, indent=2, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    sys.exit(main())
