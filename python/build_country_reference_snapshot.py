#!/usr/bin/env python3
"""Generic master workbook -> country reference-snapshot exporter
(maintainer/instructor tool).

This is the previously-missing first step of the country-scaffolding
pipeline described in docs/country_blueprint.md: turning the master
workbook's existing rows for one country into that country's tracked,
student-facing countries/<slug>/data/<slug>_reference_snapshot.csv -
the same shape as the hand-produced countries/canada/data/
canada_reference_snapshot.csv, but generic and reusable across every
assigned country instead of a one-off manual export.

It never writes to the production master workbook (read-only open,
never a save/write call against it anywhere in this file), never
infers a leader, office, party, question wording, or series alias, and
never exports a column beyond the 8 canonical ones. Country matching
uses the shared, deterministic pollitik_common.normalize_identifier()
helper - the same one that correctly finds EAD's 'Australia ' rows
(trailing-space variant) - never a reimplemented/ad hoc comparison.

The --country input is first passed through
pollitik_common.resolve_country_input() - a small, explicit alias table
for this course's assigned countries only (never fuzzy matching) that
maps a handful of known-safe spellings (e.g. "UK", "USA", "France
President") to the exact Master Country value to match against. A bare
"France" is deliberately never auto-resolved (see AmbiguousCountryError)
since Master has separate France_Pres/France_PM rows.

Usage:
    python python/build_country_reference_snapshot.py --country "Australia" --workbook data/master/pollitik_master.xlsx
"""

import argparse
import csv
import hashlib
import json
import os
import sys

import pollitik_common as pc

CANON_COLUMNS = ["Series", "Date", "Total Count", "Positive", "Neutral", "Negative", "Source", "Country"]


class SnapshotBuildError(Exception):
    """A clearly-reportable data/setup problem. Always caught in main()
    and printed as a single clean JSON error, never an unhandled
    traceback, and never a guessed/invented value."""


def sha256_of_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def default_output_path(project_dir, slug):
    return os.path.join(project_dir, "countries", slug, "data", "{}_reference_snapshot.csv".format(slug))


def _find_canonical_columns(header):
    """Maps each of the 8 canonical column names to its index in
    `header`, matching by whitespace-STRIPPED (not casefolded) exact
    header text. This handles the Master sheet's known formatting quirk
    (its 'Positive' header cell is literally '     Positive', with
    leading spaces) without guessing at, renaming, or silently matching
    any OTHER header - every other canonical name must appear verbatim
    once stripped, or this refuses loudly rather than exporting from
    the wrong column."""
    stripped_index = {}
    for i, h in enumerate(header):
        if h is None:
            continue
        key = str(h).strip()
        if key and key not in stripped_index:
            stripped_index[key] = i
    missing = [c for c in CANON_COLUMNS if c not in stripped_index]
    if missing:
        raise SnapshotBuildError(
            "Sheet is missing required column(s) (after whitespace-only "
            "header matching): {}. Headers found: {}".format(
                ", ".join(missing), [h for h in header if h not in (None, "")]
            )
        )
    return {c: stripped_index[c] for c in CANON_COLUMNS}


def _iter_canonical_rows_from_workbook(workbook_path, sheet_name):
    """Yields one dict per data row (the 8 CANON_COLUMNS, raw values
    stringified) from the master workbook. Read-only open, never a
    save/write call - mirrors the row-shaping build() always did."""
    import openpyxl

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise SnapshotBuildError(
                "Sheet '{}' not found in '{}'. Available sheets: {}".format(
                    sheet_name, workbook_path, wb.sheetnames
                )
            )
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            raise SnapshotBuildError("Sheet '{}' has no header row.".format(sheet_name))

        col = _find_canonical_columns(header)
        for row in rows_iter:
            out_row = {}
            for c in CANON_COLUMNS:
                idx = col[c]
                v = row[idx] if idx < len(row) else None
                out_row[c] = "" if v is None else str(v)
            yield out_row
    finally:
        wb.close()


def _iter_canonical_rows_from_shared_csv(csv_path):
    """Yields one dict per data row (the 8 CANON_COLUMNS) from the shared,
    tracked ALL-country CSV (python/build_reference_csv.py's output) - the
    fallback source when the master workbook itself isn't present in this
    checkout. Never writes to csv_path."""
    with open(csv_path, "r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        header = [h.strip() if h else h for h in (reader.fieldnames or [])]
        missing = [c for c in CANON_COLUMNS if c not in header]
        if missing:
            raise SnapshotBuildError(
                "Shared reference CSV '{}' is missing required column(s): "
                "{}. Headers found: {}".format(csv_path, ", ".join(missing), reader.fieldnames)
            )
        for row in reader:
            yield {c: (row.get(c) or "") for c in CANON_COLUMNS}


def build(country, workbook_path, project_dir=None, output_path=None, sheet_name="Master",
          overwrite=False, shared_reference_csv_path=None):
    project_dir = project_dir or pc.PROJECT_DIR

    try:
        resolved_country = pc.resolve_country_input(country)
        slug = pc.canonical_country_slug(country)
    except pc.AmbiguousCountryError as exc:
        raise SnapshotBuildError(str(exc))

    country_dir = os.path.join(project_dir, "countries", slug)
    if not os.path.isdir(country_dir):
        raise SnapshotBuildError(
            "countries/{}/ does not exist yet - run 'python "
            "python/init_country_workspace.py --country {}' first. This "
            "script never creates a country workspace itself.".format(slug, country)
        )

    output_path = output_path or default_output_path(project_dir, slug)
    if os.path.exists(output_path) and not overwrite:
        raise SnapshotBuildError(
            "'{}' already exists - refusing to overwrite it. Pass "
            "--overwrite only if you deliberately intend to replace the "
            "tracked snapshot.".format(output_path)
        )

    shared_reference_csv_path = shared_reference_csv_path or pc.SHARED_REFERENCE_CSV_PATH

    # Prefer the real master workbook whenever it's present (it's the
    # authoritative source); fall back to the shared, tracked all-country
    # CSV only when the workbook itself isn't available in this checkout
    # (data/master/ is git-ignored / environment-dependent) - never the
    # reverse, and never silently invented from neither.
    if os.path.isfile(workbook_path):
        source = "master_workbook"
        source_path = workbook_path
        source_sha256 = sha256_of_file(workbook_path)
        row_iter = _iter_canonical_rows_from_workbook(workbook_path, sheet_name)
    elif os.path.isfile(shared_reference_csv_path):
        source = "shared_reference_csv"
        source_path = shared_reference_csv_path
        source_sha256 = sha256_of_file(shared_reference_csv_path)
        row_iter = _iter_canonical_rows_from_shared_csv(shared_reference_csv_path)
    else:
        raise SnapshotBuildError(
            "Neither the master workbook ('{}') nor the shared reference "
            "CSV ('{}') was found. At least one is required to build a "
            "country reference snapshot - run python/build_reference_csv.py "
            "once (maintainer/instructor step, needs the real workbook) if "
            "only the shared CSV is normally missing.".format(
                workbook_path, shared_reference_csv_path
            )
        )

    target_norm = pc.normalize_identifier(resolved_country)

    matched_rows = []
    raw_variant_counts = {}
    distinct_series = set()

    for out_row in row_iter:
        raw_country = out_row.get("Country")
        if not raw_country:
            continue
        if pc.normalize_identifier(raw_country) != target_norm:
            continue

        raw_variant_counts[raw_country] = raw_variant_counts.get(raw_country, 0) + 1
        matched_rows.append(out_row)

        series_val = out_row.get("Series")
        if series_val:
            distinct_series.add(series_val)

    if not matched_rows:
        raise SnapshotBuildError(
            "Zero rows in '{}' ({}) matched Country == '{}' (resolved "
            "from input '{}', after normalization). Refusing to write "
            "an empty snapshot - nothing was written.".format(
                source_path, source, resolved_country, country
            )
        )

    out_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(out_dir, exist_ok=True)
    tmp_path = output_path + ".tmp"
    with open(tmp_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CANON_COLUMNS)
        writer.writeheader()
        for r in matched_rows:
            writer.writerow(r)
    os.replace(tmp_path, output_path)

    return {
        "country": country,
        "resolved_country": resolved_country,
        "country_slug": slug,
        "normalized_lookup_country": target_norm,
        "raw_country_variants_found": raw_variant_counts,
        "normalization_collapsed_multiple_raw_variants": len(raw_variant_counts) > 1,
        "row_count": len(matched_rows),
        "distinct_series_count": len(distinct_series),
        "output_path": output_path,
        "workbook_path": workbook_path,
        "workbook_sha256": source_sha256,
        "sheet": sheet_name,
        "overwrite_used": bool(overwrite),
        "source": source,
        "source_path": source_path,
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--country", required=True, help='e.g. "Australia"')
    parser.add_argument("--workbook", default=pc.DEFAULT_MASTER_WORKBOOK)
    parser.add_argument("--sheet", default="Master")
    parser.add_argument("--project-dir", default=None)
    parser.add_argument("--output", default=None, help="Override the tracked snapshot output path.")
    parser.add_argument("--overwrite", action="store_true", help="Deliberately overwrite an existing tracked snapshot.")
    parser.add_argument(
        "--shared-reference-csv", default=None,
        help="Override the shared all-country CSV path used as a fallback "
             "source when --workbook is not found (default: "
             "data/reference/pollitik_reference.csv).",
    )
    args = parser.parse_args(argv)

    try:
        result = build(
            args.country, args.workbook, project_dir=args.project_dir,
            output_path=args.output, sheet_name=args.sheet, overwrite=args.overwrite,
            shared_reference_csv_path=args.shared_reference_csv,
        )
    except SnapshotBuildError as exc:
        print(json.dumps({"built": False, "error": str(exc)}, indent=2, ensure_ascii=True))
        return 1

    # ensure_ascii=True (not this repo's usual False): raw_country_variants_found
    # can legitimately contain invisible Unicode (NBSP/ZWSP/etc.) - printing it
    # literally crashes on a cp1252 Windows console; \uXXXX-escaping is also more
    # legible for exactly the invisible characters this field exists to surface.
    print(json.dumps({"built": True, **result}, indent=2, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    sys.exit(main())
