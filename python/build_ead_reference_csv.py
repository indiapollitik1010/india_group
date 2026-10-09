#!/usr/bin/env python3
"""Generic EAD workbook -> shared ALL-country EAD reference CSV exporter
(maintainer/instructor tool).

Produces the tracked, student-safe
data/reference/ead_series_question_wording.csv: every country's existing
rows from "EAD Series and Question Wording.xlsx", in the same 6 canonical
columns a per-country
countries/<slug>/data/<slug>_ead_series_question_wording.csv uses (see
python/build_country_ead_snapshot.py, whose CANON_COLUMNS and
header-matching normalization this script reuses directly rather than
reimplementing). Unlike that script, this one does not filter by country -
it is the single shared source a country workspace can fall back to when
the EAD workbook itself isn't present in a checkout.

It never writes to the EAD workbook (read-only open, never a save/write
call against it anywhere in this file), never infers a missing
Description/Question Type/Question Wording/Method value, and never
exports a column beyond the 6 canonical ones.

Usage:
    python python/build_ead_reference_csv.py --ead-workbook "data/master/EAD Series and Question Wording.xlsx"
"""

import argparse
import csv
import json
import os
import sys

import pollitik_common as pc
from build_country_ead_snapshot import (
    CANON_COLUMNS,
    EadSnapshotBuildError,
    _find_canonical_ead_columns,
)
from build_country_reference_snapshot import sha256_of_file


def build(ead_workbook_path, output_path=None, sheet_name="Sheet1", overwrite=False):
    output_path = output_path or pc.SHARED_EAD_REFERENCE_CSV_PATH

    if not os.path.isfile(ead_workbook_path):
        raise EadSnapshotBuildError("EAD workbook not found at '{}'.".format(ead_workbook_path))

    if os.path.exists(output_path) and not overwrite:
        raise EadSnapshotBuildError(
            "'{}' already exists - refusing to overwrite it. Pass "
            "--overwrite only if you deliberately intend to replace the "
            "shared EAD reference CSV.".format(output_path)
        )

    workbook_sha256 = sha256_of_file(ead_workbook_path)

    import openpyxl

    wb = openpyxl.load_workbook(ead_workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise EadSnapshotBuildError(
                "Sheet '{}' not found in '{}'. Available sheets: {}".format(
                    sheet_name, ead_workbook_path, wb.sheetnames
                )
            )
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            raise EadSnapshotBuildError("Sheet '{}' has no header row.".format(sheet_name))

        col = _find_canonical_ead_columns(header, "{} sheet '{}'".format(ead_workbook_path, sheet_name))

        rows = []
        country_row_counts = {}
        for row in rows_iter:
            country_idx = col["Country"]
            raw_country = row[country_idx] if country_idx < len(row) else None
            if raw_country is None or str(raw_country).strip() == "":
                continue

            out_row = {}
            for c in CANON_COLUMNS:
                idx = col[c]
                v = row[idx] if idx < len(row) else None
                out_row[c] = "" if v is None else str(v)
            rows.append(out_row)

            country_key = pc.normalize_identifier(raw_country)
            country_row_counts[country_key] = country_row_counts.get(country_key, 0) + 1
    finally:
        wb.close()

    if not rows:
        raise EadSnapshotBuildError(
            "Zero rows with a non-empty Country found in '{}' sheet '{}'. "
            "Refusing to write an empty shared EAD reference CSV - nothing "
            "was written.".format(ead_workbook_path, sheet_name)
        )

    out_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(out_dir, exist_ok=True)
    tmp_path = output_path + ".tmp"
    with open(tmp_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CANON_COLUMNS)
        writer.writeheader()
        for r in rows:
            writer.writerow(r)
    os.replace(tmp_path, output_path)

    return {
        "row_count": len(rows),
        "country_count": len(country_row_counts),
        "country_row_counts": country_row_counts,
        "output_path": output_path,
        "output_size_bytes": os.path.getsize(output_path),
        "workbook_path": ead_workbook_path,
        "workbook_sha256": workbook_sha256,
        "sheet": sheet_name,
        "overwrite_used": bool(overwrite),
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--ead-workbook", default=pc.EAD_WORDING_REFERENCE_PATH)
    parser.add_argument("--sheet", default="Sheet1")
    parser.add_argument("--output", default=None, help="Override the tracked shared-CSV output path.")
    parser.add_argument("--overwrite", action="store_true", help="Deliberately overwrite an existing tracked shared CSV.")
    args = parser.parse_args(argv)

    try:
        result = build(args.ead_workbook, output_path=args.output, sheet_name=args.sheet, overwrite=args.overwrite)
    except EadSnapshotBuildError as exc:
        print(json.dumps({"built": False, "error": str(exc)}, indent=2, ensure_ascii=True))
        return 1

    print(json.dumps({"built": True, **result}, indent=2, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    sys.exit(main())
