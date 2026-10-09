#!/usr/bin/env python3
"""Read-only Master-vs-EAD Series coverage report for one country.

Compares a country's Master Series (read from its tracked reference
snapshot by default, or directly from the master workbook via
--use-master-workbook) against its EAD Series (from "EAD Series and
Question Wording.xlsx"), using the shared, deterministic
pollitik_common.normalize_identifier()/compare_series() helpers - the
same normalization that correctly finds EAD's whitespace-variant
Country cells (e.g. 'Australia ') instead of silently reporting zero
rows.

Never writes anything, never fabricates a missing Question Wording or
Method value, and never merges an alias/formatting-variant candidate -
those are always reported for human review only.

Usage:
    python python/report_country_ead_coverage.py --country Australia
"""

import argparse
import csv
import json
import os
import sys
from collections import Counter

import pollitik_common as pc


class ReportError(Exception):
    """A clearly-reportable setup problem. Always caught in main() and
    printed as one clean JSON error, never an unhandled traceback."""


def _master_series_counts_from_snapshot(snapshot_path):
    counts = Counter()
    with open(snapshot_path, "r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        if "Series" not in (reader.fieldnames or []):
            raise ReportError("'{}' has no 'Series' column.".format(snapshot_path))
        for row in reader:
            s = (row.get("Series") or "").strip()
            if s:
                counts[s] += 1
    return counts


def _master_series_counts_from_workbook(workbook_path, country, sheet_name="Master"):
    import openpyxl

    target_norm = pc.normalize_identifier(country)
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ReportError(
                "Sheet '{}' not found in '{}'. Available sheets: {}".format(
                    sheet_name, workbook_path, wb.sheetnames
                )
            )
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        header_stripped = [str(h).strip() if h is not None else None for h in header]
        c_idx = header_stripped.index("Country")
        s_idx = header_stripped.index("Series")
        counts = Counter()
        for r in rows:
            if pc.normalize_identifier(r[c_idx]) == target_norm and r[s_idx] is not None:
                counts[r[s_idx]] += 1
        return counts
    finally:
        wb.close()


def _ead_series_counts(ead_path, country, sheet_name="Sheet1"):
    import openpyxl

    target_norm = pc.normalize_identifier(country)
    wb = openpyxl.load_workbook(ead_path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ReportError(
                "Sheet '{}' not found in '{}'. Available sheets: {}".format(
                    sheet_name, ead_path, wb.sheetnames
                )
            )
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        header = rows[0]
        header_stripped = [str(h).strip() if h is not None else None for h in header]
        c_idx = header_stripped.index("Country")
        s_idx = header_stripped.index("Series")
        desc_idx = header_stripped.index("Description") if "Description" in header_stripped else None
        qtype_idx = header_stripped.index("Question Type") if "Question Type" in header_stripped else None
        qword_idx = header_stripped.index("Question Wording") if "Question Wording" in header_stripped else None
        method_idx = next(
            (i for i, h in enumerate(header_stripped) if h and h.startswith("Method")), None
        )

        counts = Counter()
        raw_variants = Counter()
        details = {}
        for r in rows[1:]:
            raw_country = r[c_idx]
            if raw_country is None:
                continue
            if pc.normalize_identifier(raw_country) != target_norm:
                continue
            raw_variants[raw_country] += 1
            series = r[s_idx]
            if series is None:
                continue
            counts[series] += 1
            qword = r[qword_idx] if qword_idx is not None else None
            method = r[method_idx] if method_idx is not None else None
            # Never fabricated - reports only whether the actual cell has
            # non-blank content, never invents wording/method text.
            details[series] = {
                "description": r[desc_idx] if desc_idx is not None else None,
                "question_type": r[qtype_idx] if qtype_idx is not None else None,
                "question_wording_populated": bool(qword and str(qword).strip()),
                "method_populated": bool(method and str(method).strip()),
            }
        return counts, raw_variants, details
    finally:
        wb.close()


def build_report(country, project_dir=None, snapshot_path=None, workbook_path=None,
                  ead_path=None, use_master_workbook=False):
    project_dir = project_dir or pc.PROJECT_DIR

    # Resolve aliases (e.g. "UK" -> "United Kingdom") through the one
    # centralized resolver in pollitik_common.py - never reimplemented
    # here - so both the tracked-snapshot path and every Master/EAD
    # Series match below use the same country every other pipeline
    # script would use. A deliberately-ambiguous input (bare "France")
    # is reported the same way every other setup problem in this script
    # is (ReportError), not left to propagate as a raw exception.
    try:
        resolved_country = pc.resolve_country_input(country)
        slug = pc.canonical_country_slug(country)
    except pc.AmbiguousCountryError as exc:
        raise ReportError(str(exc))

    snapshot_path = snapshot_path or os.path.join(
        project_dir, "countries", slug, "data", "{}_reference_snapshot.csv".format(slug)
    )
    workbook_path = workbook_path or pc.DEFAULT_MASTER_WORKBOOK
    ead_path = ead_path or pc.EAD_WORDING_REFERENCE_PATH

    if use_master_workbook or not os.path.isfile(snapshot_path):
        if not os.path.isfile(workbook_path):
            raise ReportError(
                "Neither the tracked snapshot ('{}') nor the master "
                "workbook ('{}') is available - cannot determine Master "
                "Series for '{}'.".format(snapshot_path, workbook_path, resolved_country)
            )
        master_counts = _master_series_counts_from_workbook(workbook_path, resolved_country)
        master_source = "master_workbook"
        master_source_path = workbook_path
    else:
        master_counts = _master_series_counts_from_snapshot(snapshot_path)
        master_source = "tracked_snapshot"
        master_source_path = snapshot_path

    if not os.path.isfile(ead_path):
        raise ReportError("EAD workbook not found at '{}'.".format(ead_path))
    ead_counts, ead_raw_variants, ead_details = _ead_series_counts(ead_path, resolved_country)

    comparison = pc.compare_series(dict(master_counts), dict(ead_counts))

    return {
        "country": country,
        "resolved_country": resolved_country,
        "country_slug": slug,
        "normalized_lookup_country": pc.normalize_identifier(resolved_country),
        "master_series_source": master_source,
        "master_series_source_path": master_source_path,
        "master_series_count": len(master_counts),
        "ead_series_count": len(ead_counts),
        "ead_raw_country_variants_found": dict(ead_raw_variants),
        "ead_series_details": ead_details,
        **comparison,
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--country", required=True)
    parser.add_argument("--snapshot", default=None, help="Override the tracked reference-snapshot CSV path.")
    parser.add_argument("--workbook", default=None, help="Override the master workbook path.")
    parser.add_argument("--ead-workbook", dest="ead_workbook", default=None, help="Override the EAD workbook path.")
    parser.add_argument(
        "--use-master-workbook", action="store_true",
        help="Force reading Master Series directly from the master workbook instead of the tracked snapshot.",
    )
    args = parser.parse_args(argv)

    try:
        result = build_report(
            args.country, snapshot_path=args.snapshot, workbook_path=args.workbook,
            ead_path=args.ead_workbook, use_master_workbook=args.use_master_workbook,
        )
    except ReportError as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, indent=2, ensure_ascii=True))
        return 1

    # ensure_ascii=True (not this repo's usual False): ead_raw_country_variants_found
    # can legitimately contain invisible Unicode - see the matching comment in
    # build_country_reference_snapshot.py.
    print(json.dumps({"ok": True, **result}, indent=2, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    sys.exit(main())
