#!/usr/bin/env python3
"""Read-only inspection of the master Pollitik workbook (Skill section 13,
"Reference Analyst" - and section 10, "the existing dataset is reference
material only").

This script never writes anything - it has no save/write call anywhere
in it. It exists so the Reference Analyst agent can learn Country
naming, series names, pollster naming, date formatting, and response
conventions from the real workbook without loading the whole file into
an LLM context window.

No workbook currently ships with this repo (data/master/ is empty until
the project owner adds the real Pollitik master workbook), so by default
this prints a clear "not found" message instead of guessing at a schema.
"""

import argparse
import json
import os
import sys

import pollitik_common as pc

DEFAULT_WORKBOOK = os.path.join(pc.MASTER_DIR, "pollitik_master.xlsx")


def inspect(workbook_path, sheet_name=None, sample_rows=5):
    if not os.path.exists(workbook_path):
        return {
            "found": False,
            "workbook_path": workbook_path,
            "message": "No workbook at this path yet. Add the real "
                       "Pollitik master workbook under data/master/ "
                       "before running the Reference Analyst for real; "
                       "nothing was inspected or invented.",
        }

    try:
        import openpyxl
    except ImportError:
        return {
            "found": True,
            "workbook_path": workbook_path,
            "error": "openpyxl is not installed. Run: pip install -r requirements.txt",
        }

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet_names = wb.sheetnames
        ws = wb[sheet_name] if sheet_name else wb[sheet_names[0]]

        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            header = []

        sample = []
        for i, row in enumerate(rows_iter):
            if i >= sample_rows:
                break
            sample.append(list(row))

        return {
            "found": True,
            "workbook_path": workbook_path,
            "sheet_names": sheet_names,
            "inspected_sheet": ws.title,
            "columns": list(header) if header else [],
            "sample_rows": sample,
        }
    finally:
        wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", default=DEFAULT_WORKBOOK)
    parser.add_argument("--sheet", default=None)
    parser.add_argument("--sample-rows", type=int, default=5)
    args = parser.parse_args()

    result = inspect(args.workbook, args.sheet, args.sample_rows)
    print(json.dumps(result, indent=2, ensure_ascii=False, default=str))
    sys.exit(0 if result.get("found") else 1)


if __name__ == "__main__":
    main()
