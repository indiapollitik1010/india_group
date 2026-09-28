"""Regression coverage for the PROJECT_DIR resolution fix in
python/pollitik_common.py.

Bug: PROJECT_DIR used to be `os.environ.get("CLAUDE_PROJECT_DIR") or
os.getcwd()`. Any pipeline script invoked from a directory other than the
repo root, with CLAUDE_PROJECT_DIR unset, silently resolved
config/allowed_domains.txt (and every other project path) against the
caller's cwd instead of the repo - making source-domain verification act
as if the allow-list were empty.

Fix: fall back to the location of pollitik_common.py itself (repo root =
two levels up from python/pollitik_common.py) instead of os.getcwd(), while
still honoring an explicitly-set CLAUDE_PROJECT_DIR.

These tests run the real python/apply_changes.py as a subprocess (as
service/agent.py does in production), always with an explicit
--staging-file and --workbook pointing at temporary fixtures created here,
and always with --dry-run, so nothing is ever read from or written to the
real data/staging/, data/master/, data/archive/, or logs/writes/.
"""

import json
import os
import subprocess
import sys

import openpyxl

from .conftest import REPO_ROOT


REAL_ALLOWED_DOMAIN_URL = "https://angusreid.org/some-report"  # real entry in config/allowed_domains.txt


def make_dry_run_record(record_id, **overrides):
    base = {
        "record_id": record_id,
        "country": "TestCountry",
        "series": "Job Approval",
        "pollster": "Test Pollster",
        "question_wording_status": "EXACT_WORDING",
        "positive": 58,
        "negative": 36,
        "neutral": None,
        "response_categories": {"Very poorly": 14, "Poorly": 22, "Well": 40, "Very well": 18},
        "category_classification": {
            "Very poorly": "negative",
            "Poorly": "negative",
            "Well": "positive",
            "Very well": "positive",
        },
        "sample_size": 1204,
        "fieldwork_date_normalized": "1/8/2026",
        "fieldwork_date_status": "OBSERVED",
        "requested_url": REAL_ALLOWED_DOMAIN_URL,
        "final_url": REAL_ALLOWED_DOMAIN_URL,
        "http_status": 200,
        "retrieved_at": "2026-08-12T14:03:00Z",
        "evidence_text": "Very poorly 14, Poorly 22, Well 40, Very well 18.",
    }
    base.update(overrides)
    return base


def make_fixture_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Record Id", "Country", "Series", "Pollster", "Positive", "Negative", "Sample Size"])
    wb.save(path)


def run_apply_dry_run(cwd, env, staging_file, workbook, record_id):
    """Runs the real python/apply_changes.py --dry-run against a
    record whose only source of validity is the *real* repo's
    config/allowed_domains.txt (angusreid.org). --staging-file and
    --workbook are always explicit temp-fixture paths, and --dry-run
    means the workbook (even though it's a throwaway fixture) is never
    actually modified."""
    proc = subprocess.run(
        [
            sys.executable, str(REPO_ROOT / "python" / "apply_changes.py"),
            "--record-id", record_id,
            "--staging-file", str(staging_file),
            "--workbook", str(workbook),
            "--sheet", "Data",
            "--dry-run",
        ],
        cwd=str(cwd),
        env=env,
        capture_output=True,
        text=True,
    )
    return proc


def test_project_dir_resolves_to_repo_from_outside_dir_env_unset(tmp_path):
    """Direct check of the fix: importing pollitik_common with cwd outside
    the repo and CLAUDE_PROJECT_DIR unset must still resolve PROJECT_DIR
    (and therefore ALLOWED_DOMAINS_PATH/MASTER_DIR/STAGING_DIR/etc.) to the
    real repo root, not to the caller's cwd."""
    outside_dir = tmp_path / "somewhere_else_entirely"
    outside_dir.mkdir()

    env = {k: v for k, v in os.environ.items() if k != "CLAUDE_PROJECT_DIR"}
    script = (
        "import sys; sys.path.insert(0, {python_dir!r}); "
        "import pollitik_common as pc; "
        "print(pc.PROJECT_DIR); print(pc.ALLOWED_DOMAINS_PATH); "
        "print(pc.MASTER_DIR); print(pc.STAGING_DIR); print(pc.ARCHIVE_DIR); "
        "print(pc.WRITES_LOG_DIR); print(pc.EAD_WORDING_REFERENCE_PATH)"
    ).format(python_dir=str(REPO_ROOT / "python"))

    proc = subprocess.run(
        [sys.executable, "-c", script],
        cwd=str(outside_dir),
        env=env,
        capture_output=True,
        text=True,
    )
    assert proc.returncode == 0, proc.stderr
    lines = proc.stdout.strip().splitlines()
    project_dir, allowed_domains_path, master_dir, staging_dir, archive_dir, writes_log_dir, ead_path = lines

    assert project_dir == str(REPO_ROOT)
    assert allowed_domains_path == str(REPO_ROOT / "config" / "allowed_domains.txt")
    assert master_dir == str(REPO_ROOT / "data" / "master")
    assert staging_dir == str(REPO_ROOT / "data" / "staging")
    assert archive_dir == str(REPO_ROOT / "data" / "archive")
    assert writes_log_dir == str(REPO_ROOT / "logs" / "writes")
    assert ead_path == str(REPO_ROOT / "docs" / "ead" / "reference" / "EAD Series and Question Wording.xlsx")

    # None of these resolved under the outside dir (the old, buggy behavior).
    for path in (allowed_domains_path, master_dir, staging_dir, archive_dir, writes_log_dir):
        assert not path.startswith(str(outside_dir))


def test_apply_changes_from_outside_repo_env_unset_finds_real_allow_list(tmp_path):
    """The end-to-end regression case: run the real apply_changes.py from a
    cwd well outside the repo, with CLAUDE_PROJECT_DIR unset. The staged
    record's only source of validity is that angusreid.org is a real entry
    in the repo's config/allowed_domains.txt - if PROJECT_DIR resolution
    fell back to cwd (the old bug), the allow-list would appear empty and
    this would be rejected instead of dry-run-succeeding."""
    outside_dir = tmp_path / "outside_repo_cwd"
    outside_dir.mkdir()

    staging_file = outside_dir / "candidates.jsonl"
    record = make_dry_run_record("regression-project-root")
    staging_file.write_text(json.dumps(record) + "\n", encoding="utf-8")

    workbook = outside_dir / "fixture.xlsx"
    make_fixture_workbook(workbook)

    env = {k: v for k, v in os.environ.items() if k != "CLAUDE_PROJECT_DIR"}
    proc = run_apply_dry_run(outside_dir, env, staging_file, workbook, "regression-project-root")

    results = json.loads(proc.stdout)
    result = results[0]
    assert "error" not in result, result
    assert result["dry_run"] is True

    # The dry run never touched the fixture workbook.
    wb = openpyxl.load_workbook(workbook, read_only=True)
    assert wb["Data"].max_row == 1


def test_apply_changes_from_repo_root_env_unset_still_works(tmp_path):
    """Normal execution from the repo root, CLAUDE_PROJECT_DIR unset, must
    keep working exactly as before the fix."""
    staging_file = tmp_path / "candidates.jsonl"
    record = make_dry_run_record("regression-project-root-repo-cwd")
    staging_file.write_text(json.dumps(record) + "\n", encoding="utf-8")

    workbook = tmp_path / "fixture.xlsx"
    make_fixture_workbook(workbook)

    env = {k: v for k, v in os.environ.items() if k != "CLAUDE_PROJECT_DIR"}
    proc = run_apply_dry_run(REPO_ROOT, env, staging_file, workbook, "regression-project-root-repo-cwd")

    results = json.loads(proc.stdout)
    result = results[0]
    assert "error" not in result, result
    assert result["dry_run"] is True


def test_explicit_claude_project_dir_still_overrides_from_outside_repo_cwd(tmp_path, sandbox):
    """An explicitly-set CLAUDE_PROJECT_DIR must still win over both the
    caller's cwd and the repo-location fallback - the existing pytest
    sandbox (tests/conftest.py) depends on this override behavior."""
    outside_dir = tmp_path / "outside_repo_cwd_2"
    outside_dir.mkdir()

    # angusreid.org (real repo allow-list) is NOT in the sandbox's
    # allow-list, so this must be rejected when CLAUDE_PROJECT_DIR points
    # at the sandbox, proving the sandbox path - not the repo - was used.
    staging_file = sandbox / "candidates.jsonl"
    record = make_dry_run_record("regression-project-root-override")
    staging_file.write_text(json.dumps(record) + "\n", encoding="utf-8")

    workbook = sandbox / "fixture.xlsx"
    make_fixture_workbook(workbook)

    proc = subprocess.run(
        [
            sys.executable, str(REPO_ROOT / "python" / "apply_changes.py"),
            "--record-id", "regression-project-root-override",
            "--staging-file", str(staging_file),
            "--workbook", str(workbook),
            "--sheet", "Data",
            "--dry-run",
        ],
        cwd=str(outside_dir),
        env={**os.environ, "CLAUDE_PROJECT_DIR": str(sandbox)},
        capture_output=True,
        text=True,
    )
    results = json.loads(proc.stdout)
    result = results[0]
    assert "error" in result
    assert "REJECTED" in result["error"]

    # And the sandbox's own approved domain still validates fine under the
    # same override, confirming this is allow-list resolution, not a
    # workbook/staging-path problem.
    record2 = make_dry_run_record(
        "regression-project-root-override-2",
        requested_url="https://approved-pollster.example/report",
        final_url="https://approved-pollster.example/report",
    )
    staging_file.write_text(json.dumps(record2) + "\n", encoding="utf-8")
    proc2 = subprocess.run(
        [
            sys.executable, str(REPO_ROOT / "python" / "apply_changes.py"),
            "--record-id", "regression-project-root-override-2",
            "--staging-file", str(staging_file),
            "--workbook", str(workbook),
            "--sheet", "Data",
            "--dry-run",
        ],
        cwd=str(outside_dir),
        env={**os.environ, "CLAUDE_PROJECT_DIR": str(sandbox)},
        capture_output=True,
        text=True,
    )
    results2 = json.loads(proc2.stdout)
    result2 = results2[0]
    assert "error" not in result2, result2
