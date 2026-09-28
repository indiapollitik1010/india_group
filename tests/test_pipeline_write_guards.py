"""Tests 7-12 and 14 from the remote-deployment task: the write-guard
chain (verify_source -> validate_record -> stage_changes -> apply_changes)
and the PreToolUse hook, exercised as real subprocesses against isolated
sandbox directories. No real Pollitik data is touched - see
tests/conftest.py for how CLAUDE_PROJECT_DIR isolation works.
"""

import datetime
import json

import openpyxl
import pytest

from .conftest import REPO_ROOT, run_hook, run_script


def make_record(**overrides):
    base = {
        "country": "TestCountry",
        "series": "Job Approval",
        "pollster": "Test Pollster",
        "question_wording_status": "EXACT_WORDING",
        "positive": 58,
        "negative": 36,
        "neutral": None,
        "response_categories": {"Very poorly": 14, "Poorly": 22, "Well": 40, "Very well": 18},
        "category_classification": {
            "Very poorly": "negative",
            "Poorly": "negative",
            "Well": "positive",
            "Very well": "positive",
        },
        "sample_size": 1204,
        "fieldwork_date_normalized": "1/8/2026",
        "fieldwork_date_status": "OBSERVED",
        "requested_url": "https://approved-pollster.example/report",
        "final_url": "https://approved-pollster.example/report",
        "http_status": 200,
        "retrieved_at": "2026-08-12T14:03:00Z",
        "evidence_text": "Very poorly 14, Poorly 22, Well 40, Very well 18.",
    }
    base.update(overrides)
    return base


def stage(sandbox, record):
    proc = run_script("stage_changes.py", sandbox, input_json=record)
    assert proc.returncode in (0, 1), proc.stderr
    return json.loads(proc.stdout)


def apply_(sandbox, record_id, workbook_path, sheet="Data"):
    proc = run_script(
        "apply_changes.py",
        sandbox,
        args=["--record-id", record_id, "--workbook", str(workbook_path), "--sheet", sheet],
    )
    results = json.loads(proc.stdout)
    return results[0]


def make_fixture_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Record Id", "Country", "Series", "Pollster", "Positive", "Negative", "Sample Size"])
    wb.save(path)


# --- Test 7: REVIEW record cannot be written -------------------------------

def test_review_record_cannot_be_written(sandbox):
    record = make_record(sample_size=-5)  # triggers SAMPLE_SIZE_REVIEW_REQUIRED
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "REVIEW"

    workbook = sandbox / "fixture.xlsx"
    make_fixture_workbook(workbook)
    result = apply_(sandbox, staged["record_id"], workbook)

    assert "error" in result
    assert "REVIEW" in result["error"]
    # Nothing beyond the header/should-be-untouched row was written.
    wb = openpyxl.load_workbook(workbook, read_only=True)
    assert wb["Data"].max_row == 1


# --- Test 8: REJECTED record cannot be written -----------------------------

def test_rejected_record_cannot_be_written(sandbox):
    record = make_record()
    del record["country"]  # triggers MISSING_REQUIRED_FIELDS -> REJECTED
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "REJECTED"

    workbook = sandbox / "fixture.xlsx"
    make_fixture_workbook(workbook)
    result = apply_(sandbox, staged["record_id"], workbook)

    assert "error" in result
    assert "REJECTED" in result["error"]


# --- Test 9: unverified source cannot be written ---------------------------

def test_unverified_source_cannot_be_written(sandbox):
    record = make_record(
        requested_url="", final_url="", http_status=None, retrieved_at="", evidence_text=""
    )
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "REJECTED"
    assert staged["source_verified"] is False
    assert "UNVERIFIED_SOURCE_URL" in staged["validation_reason"]

    workbook = sandbox / "fixture.xlsx"
    make_fixture_workbook(workbook)
    result = apply_(sandbox, staged["record_id"], workbook)
    assert "error" in result


# --- Test 10: fake/model-only URL cannot be written ------------------------

def test_model_only_url_cannot_be_written(sandbox):
    """A URL that only exists in narrative/untrusted text (never in
    final_url, i.e. never actually retrieved) must not count as
    provenance - verify_source.py only ever trusts final_url."""
    record = make_record(
        requested_url="",
        final_url="",
        http_status=None,
        retrieved_at="",
        evidence_text="",
        note="According to https://approved-pollster.example/made-up-page the president is at 58%.",
    )
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "REJECTED"
    assert staged["source_verified"] is False
    assert "UNVERIFIED_SOURCE_URL" in staged["validation_reason"]

    workbook = sandbox / "fixture.xlsx"
    make_fixture_workbook(workbook)
    result = apply_(sandbox, staged["record_id"], workbook)
    assert "error" in result


# --- Test 11: disallowed domain remains rejected ---------------------------

def test_disallowed_domain_remains_rejected(sandbox):
    record = make_record(
        requested_url="https://random-blog.example/opinion",
        final_url="https://random-blog.example/opinion",
    )
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "REJECTED"
    assert staged["source_verified"] is False
    assert "random-blog.example" in staged["validation_reason"]

    workbook = sandbox / "fixture.xlsx"
    make_fixture_workbook(workbook)
    result = apply_(sandbox, staged["record_id"], workbook)
    assert "error" in result


# --- Test 12: allowed-domain fixture succeeds ------------------------------

def test_allowed_domain_fixture_succeeds(sandbox):
    record = make_record()  # final_url on approved-pollster.example, matches sandbox allow-list
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "fixture.xlsx"
    make_fixture_workbook(workbook)
    result = apply_(sandbox, staged["record_id"], workbook)

    assert "error" not in result, result
    assert result["verified_after_save"] is True
    assert result["backup_path"] is not None
    # The backup lives under the sandbox, never the real repo.
    assert str(sandbox) in result["backup_path"]

    wb = openpyxl.load_workbook(workbook, read_only=True)
    assert wb["Data"].max_row == 2


# --- Test 14: master workbook protections remain in force ------------------

def test_hook_still_blocks_direct_production_write():
    decision = run_hook(
        "Write",
        {"file_path": "data/master/pollitik_master.xlsx", "content": "junk"},
    )
    assert decision.get("hookSpecificOutput", {}).get("permissionDecision") == "deny"


def test_hook_still_blocks_destructive_bash_on_production_workbook():
    decision = run_hook(
        "Bash",
        {"command": "rm -rf data/pollitik_master.xlsx"},
    )
    assert decision.get("hookSpecificOutput", {}).get("permissionDecision") == "deny"


def test_hook_still_allows_staging_writes():
    decision = run_hook(
        "Write",
        {"file_path": "data/staging/candidates.jsonl", "content": "{}"},
    )
    assert decision == {}  # no opinion -> allowed


def test_apply_changes_script_still_the_only_approved_writer():
    guard_source = (REPO_ROOT / ".claude" / "hooks" / "pollitik_guard.py").read_text()
    assert '"python/apply_changes.py"' in guard_source


# --- EAD manual conventions: net / app_app_dis (Skill section 46) ---------

def test_net_and_app_app_dis_computed_correctly(sandbox):
    record = make_record()  # positive=58, negative=36
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"
    assert staged["net"] == 22  # 58 - 36
    assert staged["app_app_dis"] == pytest.approx(58 / 94 * 100)


def test_net_and_app_app_dis_blank_when_negative_missing(sandbox):
    record = make_record(negative=None)
    staged = stage(sandbox, record)
    assert staged["net"] is None
    assert staged["app_app_dis"] is None


# --- EAD manual conventions: sample-size inference (Skill section 45) -----

def test_sample_size_inferred_from_pattern_routes_to_review(sandbox):
    record = make_record(sample_size_status="INFERRED_FROM_FIRM_PATTERN")
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "REVIEW"
    assert "SAMPLE_SIZE_INFERRED_FROM_PATTERN" in staged["validation_flags"]


# --- Field-mapping fix: fieldwork_date_normalized/sample_size/source_url --
# regression coverage for the two defects found by the manual copy-of-master
# write test (mapping gap + --staging-file isolation gap).

def make_master_fixture_workbook(path):
    """Mirrors the real Master worksheet's column vocabulary (Series, Date,
    Total Count, Positive, Neutral, Negative, Source, Country) rather than
    the generic "Data" fixture above, so these tests exercise the actual
    alias mapping against realistic column names."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(["Series", "Date", "Total Count", "Positive", "Neutral", "Negative", "Source", "Country"])
    wb.save(path)


def apply_master(sandbox, record_id, workbook_path, staging_file=None):
    args = ["--record-id", record_id, "--workbook", str(workbook_path), "--sheet", "Master"]
    if staging_file is not None:
        args += ["--staging-file", str(staging_file)]
    proc = run_script("apply_changes.py", sandbox, args=args)
    results = json.loads(proc.stdout)
    return results[0]


def test_fieldwork_date_and_sample_size_and_source_url_map_to_master_columns(sandbox):
    record = make_record(
        country="Canada",
        series="ANGUSREID",
        sample_size=1646,
        fieldwork_date_normalized="4/17/2026",
        source_url="https://approved-pollster.example/report",
    )
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)
    result = apply_master(sandbox, staged["record_id"], workbook)

    assert "error" not in result, result
    # Newly-aliased fields must land in mapped_fields, not the dropped bucket.
    for field in ("fieldwork_date_normalized", "sample_size", "source_url"):
        assert field in result["mapped_fields"], result
        assert field not in result["unmapped_fields_not_written_to_excel"], result

    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["Master"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    written = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
    by_col = dict(zip(header, written))

    # Date must be a native datetime, matching every existing Master row -
    # see test_fieldwork_date_written_as_native_excel_date below for the
    # dedicated regression coverage of this conversion.
    assert by_col["Date"] == datetime.datetime(2026, 4, 17)
    assert isinstance(by_col["Date"], datetime.datetime)
    assert by_col["Total Count"] == 1646
    assert by_col["Source"] == "https://approved-pollster.example/report"


def test_existing_series_positive_neutral_negative_country_mappings_preserved(sandbox):
    # positive/negative left at make_record()'s defaults (58/36) so they stay
    # arithmetically consistent with its response_categories/classification -
    # only country/series are overridden, since this test targets mapping,
    # not arithmetic validation.
    record = make_record(country="Canada", series="ANGUSREID")
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)
    result = apply_master(sandbox, staged["record_id"], workbook)

    assert "error" not in result, result
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["Master"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    written = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
    by_col = dict(zip(header, written))

    assert by_col["Series"] == "ANGUSREID"
    assert by_col["Positive"] == 58
    assert by_col["Negative"] == 36
    assert by_col["Neutral"] is None
    assert by_col["Country"] == "Canada"


def test_pm_pollster_wording_remain_unmapped_no_invented_columns(sandbox):
    """PM name, pollster, and question wording have no Master column and
    must stay in unmapped_fields_not_written_to_excel - the fix must not
    invent new workbook columns for them."""
    record = make_record(
        country="Canada", series="ANGUSREID",
        pollster="Angus Reid", executive_name="Mark Carney",
    )
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)
    result = apply_master(sandbox, staged["record_id"], workbook)

    assert "error" not in result, result
    assert "pollster" in result["unmapped_fields_not_written_to_excel"]
    assert "executive_name" in result["unmapped_fields_not_written_to_excel"]
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    header = next(wb["Master"].iter_rows(min_row=1, max_row=1, values_only=True))
    assert "Pollster" not in header
    assert "PM" not in header
    assert "Question Wording" not in header


# --- Fieldwork date native-datetime fix -------------------------------------
# Regression coverage for the date-formatting defect found by the manual
# copy-of-master integration test: fieldwork_date_normalized ("4/17/2026")
# must be written to the Master "Date" column as a native datetime, matching
# every existing Master row, not as a string.

def test_fieldwork_date_written_as_native_excel_date(sandbox):
    record = make_record(
        country="Canada", series="ANGUSREID",
        fieldwork_date_normalized="4/17/2026",
    )
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)
    result = apply_master(sandbox, staged["record_id"], workbook)
    assert "error" not in result, result

    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["Master"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    written = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
    by_col = dict(zip(header, written))

    date_value = by_col["Date"]
    assert isinstance(date_value, datetime.datetime), (
        "Date cell must be a native datetime, not a string: got %r" % (date_value,)
    )
    assert not isinstance(date_value, str)
    assert date_value == datetime.datetime(2026, 4, 17)


def test_fieldwork_date_single_digit_month_day_parsed_correctly(sandbox):
    """Skill section 19 allows either m/d/yyyy or mm/dd/yyyy - single-digit
    month/day must parse to the exact same calendar date, not be misread
    (e.g. day/month swapped) or rejected."""
    record = make_record(
        country="Canada", series="ANGUSREID",
        fieldwork_date_normalized="1/8/2026",
    )
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)
    result = apply_master(sandbox, staged["record_id"], workbook)
    assert "error" not in result, result

    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["Master"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    written = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
    by_col = dict(zip(header, written))

    assert by_col["Date"] == datetime.datetime(2026, 1, 8)  # January 8, not August 1


def test_fieldwork_date_invalid_string_preserved_not_invented(sandbox):
    """A normalized date value that does not parse as m/d/yyyy must never be
    silently coerced into a fabricated date - it is written through exactly
    as staged (unchanged) so the bad value stays visible rather than being
    guessed at."""
    record = make_record(
        country="Canada", series="ANGUSREID",
        fieldwork_date_normalized="not-a-real-date",
        fieldwork_date_status="OBSERVED",
    )
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"  # validate_record only checks presence, not format

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)
    result = apply_master(sandbox, staged["record_id"], workbook)
    assert "error" not in result, result

    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["Master"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    written = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
    by_col = dict(zip(header, written))

    assert by_col["Date"] == "not-a-real-date"


# --- Backup/log isolation: --archive-dir / --writes-log overrides ---------
# Regression coverage for the backup/log isolation defect found by the
# manual copy-of-master integration test: pointing --workbook and
# --staging-file at temporary/test locations did not stop apply_changes.py
# from writing its workbook backup into the real project's data/archive/
# and appending its write-log entry to the real project's
# logs/writes/writes.jsonl. --archive-dir/--writes-log must be honored when
# supplied, and production behavior must be unchanged when they are not.

def test_alternate_archive_dir_receives_backup_default_does_not(sandbox):
    record = make_record(country="Canada", series="ANGUSREID")
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)

    alt_archive = sandbox / "alt_archive"
    proc = run_script(
        "apply_changes.py",
        sandbox,
        args=[
            "--record-id", staged["record_id"],
            "--workbook", str(workbook),
            "--sheet", "Master",
            "--archive-dir", str(alt_archive),
        ],
    )
    result = json.loads(proc.stdout)[0]
    assert "error" not in result, result

    default_archive = sandbox / "data" / "archive"
    assert list(alt_archive.glob("*_backup_*.xlsx")), "backup missing from alternate archive dir"
    assert not list(default_archive.glob("*_backup_*.xlsx")), "backup leaked into default archive dir"
    assert result["backup_path"].startswith(str(alt_archive))


def test_alternate_writes_log_receives_entry_default_does_not(sandbox):
    record = make_record(country="Canada", series="ANGUSREID")
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)

    alt_writes_log = sandbox / "alt_logs" / "writes.jsonl"
    proc = run_script(
        "apply_changes.py",
        sandbox,
        args=[
            "--record-id", staged["record_id"],
            "--workbook", str(workbook),
            "--sheet", "Master",
            "--writes-log", str(alt_writes_log),
        ],
    )
    result = json.loads(proc.stdout)[0]
    assert "error" not in result, result

    default_writes_log = sandbox / "logs" / "writes" / "writes.jsonl"
    assert alt_writes_log.exists(), "write-log entry missing from alternate writes-log path"
    assert not default_writes_log.exists(), "write-log entry leaked into default writes.jsonl"


def test_default_archive_and_writes_log_used_when_no_overrides_passed(sandbox):
    record = make_record(country="Canada", series="ANGUSREID")
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)
    result = apply_master(sandbox, staged["record_id"], workbook)
    assert "error" not in result, result

    default_archive = sandbox / "data" / "archive"
    default_writes_log = sandbox / "logs" / "writes" / "writes.jsonl"
    assert list(default_archive.glob("*_backup_*.xlsx")), "backup missing from default archive dir"
    assert default_writes_log.exists(), "write-log entry missing from default writes.jsonl"
    assert result["backup_path"].startswith(str(default_archive))


def test_staging_file_isolation_still_works_with_archive_and_writes_log_overrides(sandbox):
    """The pre-existing --staging-file isolation fix (mark_staged_record_applied
    only ever touches the caller's --staging-file) must still hold now that
    --archive-dir/--writes-log overrides also exist alongside it."""
    record = make_record(country="Canada", series="ANGUSREID")
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    alt_staging = sandbox / "alt_staging.jsonl"
    alt_staging.write_text(json.dumps(staged) + "\n")

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)

    alt_archive = sandbox / "alt_archive2"
    alt_writes_log = sandbox / "alt_logs2" / "writes.jsonl"
    proc = run_script(
        "apply_changes.py",
        sandbox,
        args=[
            "--record-id", staged["record_id"],
            "--workbook", str(workbook),
            "--sheet", "Master",
            "--staging-file", str(alt_staging),
            "--archive-dir", str(alt_archive),
            "--writes-log", str(alt_writes_log),
        ],
    )
    result = json.loads(proc.stdout)[0]
    assert "error" not in result, result

    alt_records = [json.loads(line) for line in alt_staging.read_text().splitlines()]
    assert alt_records[0]["applied"] is True

    default_staging = sandbox / "data" / "staging" / "candidates.jsonl"
    default_records = [json.loads(line) for line in default_staging.read_text().splitlines()]
    assert all(not r.get("applied") for r in default_records)

    assert list(alt_archive.glob("*_backup_*.xlsx")), "backup missing from alternate archive dir"
    assert alt_writes_log.exists(), "write-log entry missing from alternate writes-log path"


def test_fieldwork_date_missing_written_as_none_not_invented(sandbox):
    """A record that reaches APPROVED with no normalized date at all must
    write a blank Date cell, never a fabricated one."""
    record = make_record(
        country="Canada", series="ANGUSREID",
        fieldwork_date_normalized=None,
        fieldwork_date_status="OBSERVED",
    )
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)
    result = apply_master(sandbox, staged["record_id"], workbook)
    assert "error" not in result, result

    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["Master"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    written = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
    by_col = dict(zip(header, written))

    assert by_col["Date"] is None


def test_fieldwork_date_fix_does_not_disturb_other_field_mappings(sandbox):
    """The date-type fix must be scoped to the Date column only - every
    other existing field mapping (series/positive/negative/neutral/country/
    sample_size/source_url) must still write exactly as before."""
    record = make_record(
        country="Canada", series="ANGUSREID",
        sample_size=1646,
        fieldwork_date_normalized="4/17/2026",
        source_url="https://approved-pollster.example/report",
    )
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)
    result = apply_master(sandbox, staged["record_id"], workbook)
    assert "error" not in result, result

    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["Master"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    written = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
    by_col = dict(zip(header, written))

    assert by_col["Series"] == "ANGUSREID"
    assert by_col["Total Count"] == 1646
    assert by_col["Positive"] == 58
    assert by_col["Negative"] == 36
    assert by_col["Neutral"] is None
    assert by_col["Source"] == "https://approved-pollster.example/report"
    assert by_col["Country"] == "Canada"
    assert by_col["Date"] == datetime.datetime(2026, 4, 17)


# --- --staging-file isolation fix ------------------------------------------

def test_alternate_staging_file_is_marked_applied_not_the_default_one(sandbox):
    """Reproduces the exact defect found by the manual copy-of-master write
    test: mark_staged_record_applied() must honor --staging-file end to
    end, never fall back to writing pc.STAGING_FILE (the default/real
    path) as a side effect."""
    record = make_record(country="Canada", series="ANGUSREID")
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    default_staging_file = sandbox / "data" / "staging" / "candidates.jsonl"
    assert default_staging_file.exists()
    default_before = default_staging_file.read_text(encoding="utf-8")

    # A separate, alternate staging file containing only this one record -
    # simulates the "filtered temp staging copy" pattern from the manual test.
    alt_staging_file = sandbox / "alt_candidates.jsonl"
    alt_staging_file.write_text(json.dumps(staged) + "\n", encoding="utf-8")

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)
    result = apply_master(sandbox, staged["record_id"], workbook, staging_file=alt_staging_file)

    assert "error" not in result, result

    # Only the alternate staging file gained applied=true...
    alt_lines = [json.loads(line) for line in alt_staging_file.read_text(encoding="utf-8").splitlines() if line]
    assert len(alt_lines) == 1
    assert alt_lines[0]["applied"] is True
    assert alt_lines[0]["record_id"] == staged["record_id"]

    # ...and the default/real staging file is byte-for-byte unchanged.
    assert default_staging_file.read_text(encoding="utf-8") == default_before
    default_records = [json.loads(line) for line in default_before.splitlines() if line]
    assert all("applied" not in r for r in default_records)


def test_default_staging_file_still_marked_applied_when_no_override_given(sandbox):
    """The default path (no --staging-file override) must keep working -
    the fix must not break the existing single-staging-file workflow."""
    record = make_record(country="Canada", series="ANGUSREID")
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "master_fixture.xlsx"
    make_master_fixture_workbook(workbook)
    result = apply_master(sandbox, staged["record_id"], workbook)  # no staging_file override
    assert "error" not in result, result

    default_staging_file = sandbox / "data" / "staging" / "candidates.jsonl"
    records = [json.loads(line) for line in default_staging_file.read_text(encoding="utf-8").splitlines() if line]
    matching = [r for r in records if r["record_id"] == staged["record_id"]]
    assert len(matching) == 1
    assert matching[0]["applied"] is True
