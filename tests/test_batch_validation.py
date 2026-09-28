"""Regression tests for batch validation/staging support added to
python/validate_record.py and python/stage_changes.py (the "smallest
token-efficiency improvement" from the Canada PM token-usage analysis:
let the validator agent process a whole batch of candidates in one
Python invocation instead of one process per record).

Real subprocess invocations against isolated sandbox directories - no
real Pollitik data touched. See tests/conftest.py for how
CLAUDE_PROJECT_DIR isolation works.
"""

import json

import openpyxl

from .conftest import run_hook, run_script


def make_record(**overrides):
    base = {
        "country": "TestCountry",
        "series": "Job Approval",
        "pollster": "Test Pollster",
        "question_wording_status": "EXACT_WORDING",
        "positive": 58,
        "negative": 36,
        "neutral": None,
        "response_categories": {"Very poorly": 14, "Poorly": 22, "Well": 40, "Very well": 18},
        "category_classification": {
            "Very poorly": "negative",
            "Poorly": "negative",
            "Well": "positive",
            "Very well": "positive",
        },
        "sample_size": 1204,
        "fieldwork_date_normalized": "1/8/2026",
        "fieldwork_date_status": "OBSERVED",
        "requested_url": "https://approved-pollster.example/report",
        "final_url": "https://approved-pollster.example/report",
        "http_status": 200,
        "retrieved_at": "2026-08-12T14:03:00Z",
        "evidence_text": "Very poorly 14, Poorly 22, Well 40, Very well 18.",
    }
    base.update(overrides)
    return base


def make_fixture_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Record Id", "Country", "Series", "Pollster", "Positive", "Negative", "Sample Size"])
    wb.save(path)


def apply_(sandbox, record_id, workbook_path, sheet="Data"):
    proc = run_script(
        "apply_changes.py",
        sandbox,
        args=["--record-id", record_id, "--workbook", str(workbook_path), "--sheet", sheet],
    )
    results = json.loads(proc.stdout)
    return results[0]


# --- Single-record behavior is preserved exactly ---------------------------

def test_validate_record_single_record_still_works(sandbox):
    record = make_record()
    proc = run_script("validate_record.py", sandbox, input_json=record)
    assert proc.returncode == 0, proc.stderr
    result = json.loads(proc.stdout)
    assert result["validation_status"] == "APPROVED"
    assert "batch_index" not in result  # single-record output shape is unchanged


def test_stage_changes_single_record_still_works(sandbox):
    record = make_record()
    proc = run_script("stage_changes.py", sandbox, input_json=record)
    assert proc.returncode == 0, proc.stderr
    staged = json.loads(proc.stdout)
    assert staged["validation_status"] == "APPROVED"
    assert "record_id" in staged and "staged_at" in staged
    assert "batch_index" not in staged


def test_validate_record_single_record_via_record_flag_still_works(sandbox):
    record_path = sandbox / "single.json"
    record_path.write_text(json.dumps(make_record()))
    proc = run_script("validate_record.py", sandbox, args=["--record", str(record_path)])
    assert proc.returncode == 0, proc.stderr
    result = json.loads(proc.stdout)
    assert result["validation_status"] == "APPROVED"


# --- Batch validation: JSON array via --records, JSONL via --batch/stdin ---

def test_validate_record_batch_via_records_file_json_array(sandbox):
    records = [make_record(pollster="Firm A"), make_record(pollster="Firm B")]
    records_path = sandbox / "batch_in.json"
    records_path.write_text(json.dumps(records))

    proc = run_script("validate_record.py", sandbox, args=["--records", str(records_path)])
    assert proc.returncode == 0, proc.stderr
    results = json.loads(proc.stdout)
    assert len(results) == 2
    assert [r["batch_index"] for r in results] == [0, 1]
    assert all(r["validation_status"] == "APPROVED" for r in results)
    assert results[0]["pollster"] == "Firm A"
    assert results[1]["pollster"] == "Firm B"


def test_validate_record_batch_via_stdin_jsonl(sandbox):
    records = [make_record(pollster="Firm A"), make_record(pollster="Firm B")]
    jsonl_text = "\n".join(json.dumps(r) for r in records)

    proc = run_script("validate_record.py", sandbox, args=["--batch"], raw_input=jsonl_text)
    assert proc.returncode == 0, proc.stderr
    results = json.loads(proc.stdout)
    assert len(results) == 2
    assert all(r["validation_status"] == "APPROVED" for r in results)


# --- Batch staging -----------------------------------------------------------

def test_stage_changes_batch_via_records_file(sandbox):
    records = [make_record(pollster="Firm A"), make_record(pollster="Firm B")]
    records_path = sandbox / "batch_in.json"
    records_path.write_text(json.dumps(records))

    proc = run_script("stage_changes.py", sandbox, args=["--records", str(records_path)])
    assert proc.returncode == 0, proc.stderr
    staged_list = json.loads(proc.stdout)
    assert len(staged_list) == 2
    record_ids = {s["record_id"] for s in staged_list}
    assert len(record_ids) == 2  # each got a distinct id

    staging_file = sandbox / "data" / "staging" / "candidates.jsonl"
    lines = [json.loads(line) for line in staging_file.read_text().splitlines() if line.strip()]
    staged_ids_on_disk = {rec["record_id"] for rec in lines}
    assert record_ids.issubset(staged_ids_on_disk)  # both records actually persisted


# --- Mixed statuses in one batch --------------------------------------------

def test_batch_handles_mixed_statuses_correctly(sandbox):
    approved = make_record(pollster="Approved Firm")
    review = make_record(pollster="Review Firm", sample_size=-5)  # -> SAMPLE_SIZE_REVIEW_REQUIRED
    rejected = make_record(pollster="Rejected Firm")
    del rejected["country"]  # -> MISSING_REQUIRED_FIELDS

    proc = run_script(
        "stage_changes.py", sandbox, args=["--batch"],
        raw_input=json.dumps([approved, review, rejected]),
    )
    assert proc.returncode == 0, proc.stderr
    staged_list = json.loads(proc.stdout)
    assert [s["validation_status"] for s in staged_list] == ["APPROVED", "REVIEW", "REJECTED"]
    assert [s["batch_index"] for s in staged_list] == [0, 1, 2]
    # every status still gets staged, not just APPROVED (existing staging semantics)
    assert all(s.get("record_id") for s in staged_list)


# --- One invalid record does not corrupt the rest of the batch -------------

def test_bad_jsonl_line_does_not_corrupt_other_records(sandbox):
    good_a = json.dumps(make_record(pollster="Firm A"))
    good_b = json.dumps(make_record(pollster="Firm B"))
    jsonl_text = "\n".join([good_a, "{not valid json", good_b])

    proc = run_script("validate_record.py", sandbox, args=["--batch"], raw_input=jsonl_text)
    results = json.loads(proc.stdout)
    assert len(results) == 3  # all three lines produced a result - nothing dropped
    assert results[0]["validation_status"] == "APPROVED"
    assert results[0]["pollster"] == "Firm A"
    assert results[1]["validation_status"] == "REJECTED"
    assert "INVALID_INPUT_RECORD" in results[1]["validation_flags"]
    assert results[2]["validation_status"] == "APPROVED"
    assert results[2]["pollster"] == "Firm B"


def test_non_dict_array_entry_does_not_corrupt_other_records(sandbox):
    records = [make_record(pollster="Firm A"), "not a record", make_record(pollster="Firm B")]

    proc = run_script("stage_changes.py", sandbox, args=["--batch"], raw_input=json.dumps(records))
    staged_list = json.loads(proc.stdout)
    assert len(staged_list) == 3
    assert staged_list[0]["validation_status"] == "APPROVED"
    assert staged_list[1]["validation_status"] == "REJECTED"
    assert "INVALID_INPUT_RECORD" in staged_list[1]["validation_flags"]
    assert staged_list[2]["validation_status"] == "APPROVED"


def test_empty_batch_returns_empty_result_list(sandbox):
    proc = run_script("validate_record.py", sandbox, args=["--batch"], raw_input="")
    assert proc.returncode == 0, proc.stderr
    assert json.loads(proc.stdout) == []


# --- Production workbook write protections are unaffected ------------------

def test_batch_staged_approved_record_can_be_applied_review_cannot(sandbox):
    approved = make_record(pollster="Firm A")
    review = make_record(pollster="Firm B", sample_size=-5)

    proc = run_script(
        "stage_changes.py", sandbox, args=["--batch"],
        raw_input=json.dumps([approved, review]),
    )
    staged_list = json.loads(proc.stdout)
    approved_id = staged_list[0]["record_id"]
    review_id = staged_list[1]["record_id"]
    assert staged_list[0]["validation_status"] == "APPROVED"
    assert staged_list[1]["validation_status"] == "REVIEW"

    workbook = sandbox / "fixture.xlsx"
    make_fixture_workbook(workbook)

    approved_result = apply_(sandbox, approved_id, workbook)
    assert "error" not in approved_result, approved_result
    assert approved_result["verified_after_save"] is True

    review_result = apply_(sandbox, review_id, workbook)
    assert "error" in review_result
    assert "REVIEW" in review_result["error"]

    # Only the one APPROVED row was written.
    wb = openpyxl.load_workbook(workbook, read_only=True)
    assert wb["Data"].max_row == 2  # header + the one approved row


def test_hook_still_blocks_direct_production_write():
    decision = run_hook(
        "Write",
        {"file_path": "data/master/pollitik_master.xlsx", "content": "junk"},
    )
    assert decision.get("hookSpecificOutput", {}).get("permissionDecision") == "deny"


def test_hook_still_blocks_destructive_bash_on_production_workbook():
    decision = run_hook(
        "Bash",
        {"command": "rm -rf data/pollitik_master.xlsx"},
    )
    assert decision.get("hookSpecificOutput", {}).get("permissionDecision") == "deny"


def test_apply_changes_script_still_the_only_approved_writer():
    from .conftest import REPO_ROOT
    guard_source = (REPO_ROOT / ".claude" / "hooks" / "pollitik_guard.py").read_text()
    assert '"python/apply_changes.py"' in guard_source
