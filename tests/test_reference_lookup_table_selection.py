"""Regression coverage for python/reference_lookup.py's default-table
selection.

Bug this guards against: when data/reference_index.duckdb is built from
the FULL master workbook (every sheet becomes its own table), the old
code picked `table_name or tables[0]` - i.e. whichever table happened to
come back first from information_schema.tables - which could silently
land on an unrelated sheet (e.g. 'arg2023') instead of the canonical
'master' table. A student querying --country Australia against that
index got zero/wrong results with no indication anything was wrong.

These tests build small synthetic Excel/CSV fixtures and run the real
python/build_reference_index.py + python/reference_lookup.py as
subprocesses against isolated sandbox directories - never the real
repo's data/master/. A couple of integration tests at the bottom run
against the REAL authoritative workbooks and are skipped (not failed)
when those git-ignored files aren't present in this checkout.
"""

import json

import openpyxl
import pytest

from .conftest import REPO_ROOT, run_script

REAL_MASTER_WORKBOOK = REPO_ROOT / "data" / "master" / "pollitik_master.xlsx"


def make_multi_sheet_workbook_with_master(path):
    """Simulates a full-master-workbook-built index: a canonical 'Master'
    sheet (normalizes to table 'master') plus unrelated extra sheets with
    generic/unnamed-looking headers, exactly like the real workbook's
    Sheet11/Sheet22/Arg2023 auxiliary sheets."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(["Series", "Date", "Total Count", "     Positive", "Neutral", "Negative", None, None, "Source", "Country"])
    ws.append(["NEWSPOLL", "2020-01-15 00:00:00", 1000, 55, None, 40, None, None, "https://x/1", "Australia"])
    ws.append(["NEWSPOLL", "2020-02-15 00:00:00", 1000, 56, None, 39, None, None, "https://x/2", "Australia"])
    ws.append(["ANGUSREID", "2020-01-15 00:00:00", 1000, 50, None, 45, None, None, "https://x/3", "Canada"])

    ws2 = wb.create_sheet("Arg2023")
    ws2.append([None, None, None])
    ws2.append(["some", "unrelated", "data"])

    ws3 = wb.create_sheet("Sheet11")
    ws3.append(["ELABE", "2016-01-06 00:00:00", 1000])
    ws3.append(["ELABE", "2016-02-06 00:00:00", 1001])

    wb.save(path)


def make_multi_sheet_workbook_no_master(path):
    """No sheet named/normalizing to 'master' anywhere - the genuinely
    ambiguous case that must fail clearly rather than guess."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SheetA"
    ws.append(["Series", "Date", "Country"])
    ws.append(["X", "2020-01-01", "TestCountry"])

    ws2 = wb.create_sheet("SheetB")
    ws2.append(["Series", "Date", "Country"])
    ws2.append(["Y", "2020-01-01", "TestCountry"])

    wb.save(path)


def make_country_snapshot_csv(path, country="Australia"):
    path.write_text(
        "Series,Date,Total Count,Positive,Neutral,Negative,Source,Country\n"
        "NEWSPOLL,1/15/2020,1000,55,,40,https://x/1,{country}\n"
        "NEWSPOLL,2/15/2020,1000,56,,39,https://x/2,{country}\n".format(country=country),
        encoding="utf-8",
    )


def build_index(sandbox, workbook, index_path=None, sheet=None):
    index_path = index_path or (sandbox / "data" / "reference_index.duckdb")
    args = ["--workbook", str(workbook), "--index", str(index_path)]
    if sheet:
        args += ["--sheet", sheet]
    proc = run_script("build_reference_index.py", sandbox, args=args)
    assert proc.returncode == 0, proc.stderr
    return json.loads(proc.stdout), index_path


def run_lookup(sandbox, index_path, extra_args=None):
    proc = run_script(
        "reference_lookup.py", sandbox,
        args=["--index", str(index_path), *(extra_args or [])],
    )
    return proc, json.loads(proc.stdout)


# --- required regression cases -------------------------------------------

def test_full_master_index_with_other_sheets_defaults_to_master(sandbox):
    workbook = sandbox / "master.xlsx"
    make_multi_sheet_workbook_with_master(workbook)
    _, index_path = build_index(sandbox, workbook)

    proc, result = run_lookup(sandbox, index_path, ["--country", "Australia"])
    assert proc.returncode == 0
    assert result["found_index"] is True
    assert result["table"] == "master"
    assert result["table_source"] == "master_default"
    assert result["row_count_returned"] == 2
    assert all(r["Country"] == "Australia" for r in result["rows"])


def test_snapshot_built_index_single_table_works_automatically(sandbox):
    """Preserves the country-snapshot-built index path (e.g. built from
    countries/australia/data/australia_reference_snapshot.csv) - a single
    table with a country-specific name, no 'master' anywhere."""
    snapshot = sandbox / "australia_reference_snapshot.csv"
    make_country_snapshot_csv(snapshot, country="Australia")
    built, index_path = build_index(sandbox, snapshot)
    assert built["tables"][0]["table"] == "australia_reference_snapshot"

    proc, result = run_lookup(sandbox, index_path, ["--country", "Australia"])
    assert proc.returncode == 0
    assert result["found_index"] is True
    assert result["table"] == "australia_reference_snapshot"
    assert result["table_source"] == "single_table_default"
    assert result["row_count_returned"] == 2


def test_multiple_tables_no_master_fails_clearly(sandbox):
    workbook = sandbox / "no_master.xlsx"
    make_multi_sheet_workbook_no_master(workbook)
    _, index_path = build_index(sandbox, workbook)

    proc, result = run_lookup(sandbox, index_path, ["--country", "TestCountry"])
    assert proc.returncode == 1  # a clear failure, not a silent wrong-table guess
    assert result["found_index"] is True  # the index itself does exist
    assert result["ambiguous_tables"] is True
    assert result["rows"] == []
    assert set(result["available_tables"]) == {"sheeta", "sheetb"}
    assert "master" not in result["message"].lower() or "none is named" in result["message"].lower()


def test_explicit_table_still_overrides_default(sandbox):
    workbook = sandbox / "master.xlsx"
    make_multi_sheet_workbook_with_master(workbook)
    _, index_path = build_index(sandbox, workbook)

    proc, result = run_lookup(sandbox, index_path, ["--table", "sheet11"])
    assert proc.returncode == 0
    assert result["found_index"] is True
    assert result["table"] == "sheet11"
    assert result["table_source"] == "explicit"


def test_explicit_table_not_found_still_reports_clearly(sandbox):
    workbook = sandbox / "master.xlsx"
    make_multi_sheet_workbook_with_master(workbook)
    _, index_path = build_index(sandbox, workbook)

    proc, result = run_lookup(sandbox, index_path, ["--table", "does_not_exist"])
    assert result["found_index"] is True
    assert "not found" in result["message"].lower()
    assert result["rows"] == []


# --- integration tests against the REAL authoritative workbook -----------

pytestmark_real = pytest.mark.skipif(
    not REAL_MASTER_WORKBOOK.exists(),
    reason="real data/master/pollitik_master.xlsx not present in this checkout (git-ignored, environment-dependent)",
)


@pytestmark_real
def test_real_australia_lookup_works_without_specifying_table(sandbox):
    """The exact regression this fix targets: querying the real,
    full-master-built index for Australia must return real Australia
    rows without the caller ever passing --table master."""
    index_path = sandbox / "data" / "reference_index.duckdb"
    build_proc = run_script(
        "build_reference_index.py", sandbox,
        args=["--workbook", str(REAL_MASTER_WORKBOOK), "--index", str(index_path)],
    )
    assert build_proc.returncode == 0, build_proc.stderr

    proc, result = run_lookup(sandbox, index_path, ["--country", "Australia", "--series", "NEWSPOLL", "--limit", "3"])
    assert proc.returncode == 0
    assert result["found_index"] is True
    assert result["table"] == "master"
    assert result["table_source"] == "master_default"
    assert result["row_count_returned"] == 3
    assert all(r["Country"] == "Australia" for r in result["rows"])
    assert all(r["Series"] == "NEWSPOLL" for r in result["rows"])


@pytestmark_real
def test_real_canada_lookup_regression(sandbox):
    """Canada must keep working through the same default-table path -
    this fix must not regress the country it was already correct for."""
    index_path = sandbox / "data" / "reference_index.duckdb"
    build_proc = run_script(
        "build_reference_index.py", sandbox,
        args=["--workbook", str(REAL_MASTER_WORKBOOK), "--index", str(index_path)],
    )
    assert build_proc.returncode == 0, build_proc.stderr

    proc, result = run_lookup(sandbox, index_path, ["--country", "Canada", "--series", "ANGUSREID", "--limit", "3"])
    assert proc.returncode == 0
    assert result["found_index"] is True
    assert result["table"] == "master"
    assert result["table_source"] == "master_default"
    assert result["row_count_returned"] > 0
    assert all(r["Country"] == "Canada" for r in result["rows"])
