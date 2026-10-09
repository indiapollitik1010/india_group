"""Tests for this task's additions: the Sample_Size/"Total Count" alias,
optional-field (Negative/Sample_Size/Pollster) validation behavior, the
domain-approval-request workflow (check_domain.py + the hook), and a few
global-vs-country-specific invariants. Real subprocess invocations
against isolated sandbox directories, same pattern as
tests/test_pipeline_write_guards.py - no real Pollitik data touched.
"""

import glob
import json
from pathlib import Path

import openpyxl

from .conftest import REPO_ROOT, run_hook, run_script


def make_record(**overrides):
    base = {
        "country": "TestCountry",
        "series": "Job Approval",
        "pollster": "Test Pollster",
        "question_wording_status": "EXACT_WORDING",
        "positive": 58,
        "negative": 36,
        "neutral": None,
        "response_categories": {"Well": 58, "Poorly": 36},
        "category_classification": {"Well": "positive", "Poorly": "negative"},
        "sample_size": 1204,
        "fieldwork_date_normalized": "1/8/2026",
        "fieldwork_date_status": "OBSERVED",
        "requested_url": "https://approved-pollster.example/report",
        "final_url": "https://approved-pollster.example/report",
        "http_status": 200,
        "retrieved_at": "2026-08-12T14:03:00Z",
        "evidence_text": "Well 58, Poorly 36.",
    }
    base.update(overrides)
    return base


def stage(sandbox, record):
    proc = run_script("stage_changes.py", sandbox, input_json=record)
    assert proc.returncode in (0, 1), proc.stderr
    return json.loads(proc.stdout)


def apply_(sandbox, record_id, workbook_path, sheet="Data"):
    proc = run_script(
        "apply_changes.py",
        sandbox,
        args=["--record-id", record_id, "--workbook", str(workbook_path), "--sheet", sheet],
    )
    results = json.loads(proc.stdout)
    return results[0]


# --- 1 & 2: "Total Count" recognized as Sample_Size, no duplicate column --

def make_fixture_workbook_total_count(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Record Id", "Country", "Series", "Pollster", "Positive", "Negative", "Total Count"])
    wb.save(path)


def test_total_count_column_recognized_as_sample_size(sandbox):
    # apply_changes.py resolves aliases from config/schema_mapping.yaml
    # under CLAUDE_PROJECT_DIR (the sandbox here), so the sandbox needs
    # its own copy - mirrors the real repo's config/schema_mapping.yaml
    # entry for Sample_Size, scoped to just what this test exercises.
    (sandbox / "config" / "schema_mapping.yaml").write_text(
        'fields:\n'
        '  Sample_Size:\n'
        '    internal_field: sample_size\n'
        '    aliases: ["Total Count", "Total_Count"]\n'
    )

    record = make_record()  # sample_size=1204
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"

    workbook = sandbox / "fixture_total_count.xlsx"
    make_fixture_workbook_total_count(workbook)
    result = apply_(sandbox, staged["record_id"], workbook)

    assert "error" not in result, result
    assert result["mapped_fields"]["sample_size"] == 1204
    assert "sample_size" not in result["unmapped_fields_not_written_to_excel"]

    wb = openpyxl.load_workbook(workbook, read_only=True)
    ws = wb["Data"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    data_row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    total_count_idx = header.index("Total Count")
    assert data_row[total_count_idx] == 1204
    # No duplicate "Sample_Size"-ish column was created alongside it.
    assert len(header) == 7
    assert not any("sample" in str(h).lower() and h != "Total Count" for h in header)


# --- 3 & 4: missing Negative / missing Sample_Size never auto-reject -------

def test_missing_negative_does_not_reject(sandbox):
    record = make_record(
        negative=None,
        response_categories={"Well": 58},
        category_classification={"Well": "positive"},
    )
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"
    assert staged["negative"] is None


def test_missing_sample_size_does_not_reject(sandbox):
    record = make_record(sample_size=None)
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "APPROVED"
    assert staged["sample_size"] is None


# --- 5: missing Pollster does not reject when series is established -------

def test_missing_pollster_routes_to_review_not_rejected(sandbox):
    record = make_record(pollster=None)
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "REVIEW"
    assert "SERIES_REVIEW_REQUIRED" in staged["validation_flags"]
    assert staged["validation_status"] != "REJECTED"


# --- 6: uncertain pollster/series relationship -> SERIES_REVIEW_REQUIRED --

def test_uncertain_series_match_status_routes_to_review(sandbox):
    record = make_record(series_match_status="SERIES_REVIEW_REQUIRED")
    staged = stage(sandbox, record)
    assert staged["validation_status"] == "REVIEW"
    assert "SERIES_REVIEW_REQUIRED" in staged["validation_flags"]


# --- 7 & 8: unapproved domain -> DOMAIN_APPROVAL_REQUIRED, never fetched --

def test_check_domain_unapproved_domain(sandbox):
    proc = run_script("check_domain.py", sandbox, args=["--url", "https://random-blog.example/opinion"])
    assert proc.returncode == 1
    result = json.loads(proc.stdout)
    assert result["approved"] is False
    assert result["status"] == "DOMAIN_APPROVAL_REQUIRED"


def test_hook_denies_webfetch_to_unapproved_domain():
    decision = run_hook("WebFetch", {"url": "https://random-blog.example/opinion"})
    assert decision.get("hookSpecificOutput", {}).get("permissionDecision") == "deny"


# --- 9: approved domain follows the normal retrieval flow -----------------

def test_check_domain_approved_domain(sandbox):
    proc = run_script("check_domain.py", sandbox, args=["--url", "https://approved-pollster.example/report"])
    assert proc.returncode == 0
    result = json.loads(proc.stdout)
    assert result["approved"] is True
    assert result["status"] == "APPROVED"


def test_hook_allows_webfetch_to_approved_domain():
    decision = run_hook("WebFetch", {"url": "https://approved-pollster.example/report"})
    assert decision == {}  # no opinion -> allowed


# --- 10: fabricated/model-only URLs remain invalid (regression guard) -----
# (Primary coverage already lives in test_pipeline_write_guards.py:
#  test_model_only_url_cannot_be_written / test_unverified_source_cannot_be_written.)

def test_check_domain_reports_invalid_url_cleanly(sandbox):
    proc = run_script("check_domain.py", sandbox, args=["--url", "not-a-url"])
    assert proc.returncode == 1
    result = json.loads(proc.stdout)
    assert result["approved"] is False
    assert result["status"] == "INVALID_URL"


# --- 11: writer logic is not Canada-specific -------------------------------

def test_writer_and_hook_contain_no_hardcoded_country_names():
    country_tokens = ["Canada", "Germany", "France", "CANADA_", "canada_pm"]
    for path in (REPO_ROOT / "python" / "apply_changes.py", REPO_ROOT / ".claude" / "hooks" / "pollitik_guard.py"):
        source = path.read_text(encoding="utf-8")
        for token in country_tokens:
            assert token not in source, f"{path} unexpectedly references {token!r}"


# --- 12: modular domain config scaffold has no non-scoped domain data -----

def test_domain_config_scaffold_matches_allowed_domains_and_stays_scoped():
    import yaml

    template_yaml = yaml.safe_load((REPO_ROOT / "config" / "domains" / "_TEMPLATE.yaml").read_text())
    assert template_yaml["domains"] == []
    assert template_yaml["country"] == "Example Country"

    allowed = set(
        line.strip().lower()
        for line in (REPO_ROOT / "config" / "allowed_domains.txt").read_text().splitlines()
        if line.strip() and not line.strip().startswith("#")
    )

    domain_files = sorted(Path(p) for p in glob.glob(str(REPO_ROOT / "config" / "domains" / "*.yaml")))
    for path in domain_files:
        if path.name == "_TEMPLATE.yaml":
            continue
        data = yaml.safe_load(path.read_text())
        declared_country = data.get("country")  # None/absent for global.yaml
        for entry in data.get("domains") or []:
            # Metadata never grants access on its own - every domain
            # described here must already be a real enforced entry.
            assert entry["host"].lower() in allowed, (
                f"{path.name}: {entry['host']} has metadata but is not in "
                "config/allowed_domains.txt"
            )
            # Rule: country-specific domain data stays associated only
            # with its relevant country - a per-country file must not
            # describe a domain belonging to a different country.
            if declared_country is not None:
                assert entry.get("country") == declared_country, (
                    f"{path.name}: entry for {entry['host']} has country "
                    f"{entry.get('country')!r}, expected {declared_country!r}"
                )


# --- 13: global rules remain available without duplication ----------------

def test_every_agent_references_the_canonical_skill_rather_than_duplicating_it():
    for path in glob.glob(str(REPO_ROOT / ".claude" / "agents" / "*.md")):
        source = Path(path).read_text(encoding="utf-8")
        assert ".claude/skills/pollitik-executive-support/SKILL.md" in source, (
            f"{path} does not reference the canonical Skill"
        )
